from rest_framework.viewsets import ModelViewSet
from rest_framework.permissions import IsAuthenticated
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework import status
from django.http import HttpResponse
import csv
import io
import os
from datetime import date, datetime
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError


from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter, OrderingFilter
from rest_framework.parsers import JSONParser, MultiPartParser, FormParser
from rest_framework.exceptions import ValidationError as APIValidationError

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill

from .models import Employee, Department
from ACCOUNTS.models import CURRENT_USER
from .serializers import (
    EmployeeSerializer,
    DepartmentSerializer
)
from .permissions import IsAdminOrHRStrict

# ==========================================================
# Employee ViewSet
# ==========================================================
class EmployeeViewSet(ModelViewSet):

    queryset = (
        Employee.objects
        .select_related(
            "department",
            "bank",
            "extra",
            "documents",
        )
        .order_by("-created_at")
    )

    serializer_class = EmployeeSerializer

    parser_classes = (
        JSONParser,
        MultiPartParser,
        FormParser,
    )

    permission_classes = [
        IsAuthenticated,
        IsAdminOrHRStrict,
    ]

    filter_backends = [
        DjangoFilterBackend,
        SearchFilter,
        OrderingFilter,
    ]

    filterset_fields = [
        "department",
        "is_active",
        "employment_type",
        "work_location",
    ]

    search_fields = [
        "employee_id",
        "first_name",
        "middle_name",
        "last_name",
        "position",
        "official_email",
    ]

    ordering_fields = [
        "first_name",
        "last_name",
        "joining_date",
        "salary",
        "created_at",
    ]

    ordering = ["-created_at"]

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context["request"] = self.request
        return context

    @action(detail=False, methods=["get"], url_path="import-template")
    def import_template(self, request):
        headers = [
            "First Name", "Middle Name", "Last Name", "Gender", "Date of Birth",
            "Phone Number", "Official Email", "Personal Email", "Marital Status",
            "Blood Group", "Nationality", "Father Name", "Mother Name", "Position",
            "Employee Type", "Attendance Category", "Employment Type", "Department",
            "Work Location", "Joining Date", "Salary", "PAN", "Aadhaar", "UAN",
            "PF Number", "Status",
        ]
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Employees"
        worksheet.append(headers)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=len(headers)).column_letter}1"
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="2563EB")
        for column_name, values in {
            "Gender": "Male,Female,Other",
            "Marital Status": "Single,Married,Divorced,Widowed",
            "Attendance Category": "In-house,On-site",
            "Employment Type": "Full Time,Intern,Contract,Part Time",
            "Status": "Active,Inactive",
        }.items():
            column_index = headers.index(column_name) + 1
            column_letter = worksheet.cell(row=1, column=column_index).column_letter
            validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
            worksheet.add_data_validation(validation)
            validation.add(f"{column_letter}2:{column_letter}1000")
        for column_name in ("Phone Number", "PAN", "Aadhaar", "UAN", "PF Number"):
            column_letter = worksheet.cell(row=1, column=headers.index(column_name) + 1).column_letter
            for row_index in range(2, 1001):
                worksheet[f"{column_letter}{row_index}"].number_format = "@"
        for column in worksheet.columns:
            worksheet.column_dimensions[column[0].column_letter].width = min(
                max(len(str(cell.value or "")) for cell in column) + 3, 28
            )

        instructions = workbook.create_sheet("Instructions")
        instruction_rows = [
            ["Employee import guide"],
            ["Enter new employees on the Employees sheet. Keep its first row of column names unchanged."],
            ["Required: First Name, Last Name, Phone Number, Official Email, Position, Joining Date, Salary, PAN, Aadhaar."],
            ["Employment Type: Full Time, Intern, Contract, or Part Time."],
            ["Attendance Category: In-house or On-site."],
            ["Department is optional; use a department name that already exists in this HR system."],
            ["PF Number is optional and must contain no more than 20 characters. Format ID numbers as Text in Excel."],
            ["Dates: YYYY-MM-DD (example: 2026-01-31). Excel date cells are also accepted."],
            ["Status is optional: Active or Inactive. Blank defaults to active."],
            ["Employee ID is generated by the system. Do not add it to the import sheet."],
            ["The Example - Reference Only sheet is not imported. Copy its sample row into Employees only after replacing every sample value."],
        ]
        for row in instruction_rows:
            instructions.append(row)
        instructions.column_dimensions["A"].width = 110
        instructions.freeze_panes = "A2"
        for cell in instructions[1]:
            cell.font = Font(bold=True, size=15, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="2563EB")

        departments_sheet = workbook.create_sheet("Departments")
        departments_sheet.append(["Existing department names - use one of these values"])
        for department_name in Department.objects.order_by("name").values_list("name", flat=True):
            departments_sheet.append([department_name])
        if departments_sheet.max_row == 1:
            departments_sheet.append(["No departments have been created yet. Department is optional."])
        departments_sheet.column_dimensions["A"].width = 70
        departments_sheet.freeze_panes = "A2"
        departments_sheet["A1"].font = Font(bold=True, color="FFFFFF")
        departments_sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="2563EB")

        example = workbook.create_sheet("Example - Reference Only")
        example.append(headers)
        example.append([
            "Sample", "", "Employee", "Other", "1995-01-15", "9000000000",
            "sample.employee@example.com", "", "Single", "O+", "Indian", "", "",
            "Sample Analyst", "Permanent", "In-house", "Full Time",
            "Replace with an existing department", "Office", "2026-01-01", 50000,
            "ABCDE1234F", "123456789012", "", "", "Active",
        ])
        example.freeze_panes = "A2"
        example.auto_filter.ref = f"A1:{example.cell(row=1, column=len(headers)).column_letter}2"
        example.sheet_properties.tabColor = "F59E0B"
        for cell in example[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="64748B")
        for column in example.columns:
            example.column_dimensions[column[0].column_letter].width = min(
                max(len(str(cell.value or "")) for cell in column) + 3, 28
            )

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="employee_import_template.xlsx"'
        workbook.save(response)
        return response

    def perform_create(self, serializer):
        employee_email = serializer.validated_data.get("official_email", "")
        employee_login = CURRENT_USER.objects.filter(
            email__iexact=employee_email,
            role="viewer",
        ).first()
        serializer.save(user=employee_login)

    def perform_update(self, serializer):
        employee = serializer.instance
        login = employee.user
        if login is None:
            employee_email = serializer.validated_data.get(
                "official_email", employee.official_email
            )
            login = CURRENT_USER.objects.filter(
                email__iexact=employee_email,
                role="viewer",
            ).first()
        serializer.save(user=login)

    @action(detail=True, methods=["post"], url_path="link-account")
    def link_account(self, request, pk=None):
        employee = self.get_object()
        email = str(request.data.get("email") or employee.official_email).strip()
        if not email:
            raise APIValidationError({"email": "Enter the employee login email."})

        user = CURRENT_USER.objects.filter(email__iexact=email).first()
        if user is None:
            raise APIValidationError({
                "email": (
                    "No login account exists for this email. The employee must "
                    "create a viewer account using this email first."
                )
            })
        if user.role != "viewer":
            raise APIValidationError({
                "email": "Link an employee viewer account, not an HR or admin account."
            })
        if Employee.objects.filter(user=user).exclude(pk=employee.pk).exists():
            raise APIValidationError({
                "email": "This login account is already linked to another employee."
            })

        employee.user = user
        employee.save(update_fields=["user"])
        return Response({
            "message": "Employee login linked successfully.",
            "employee_id": employee.pk,
            "linked_user_email": user.email,
        })

    # ==========================================
    # EXPORT
    # ==========================================
    @action(
        detail=False,
        methods=["get"],
        url_path="export"
    )
    def export_employees(self, request):

        employees = self.get_queryset()

        ids = request.query_params.get("ids")

        if ids:
            employee_ids = [
                value.strip()
                for value in ids.split(",")
                if value.strip()
            ]

            employees = employees.filter(
                id__in=employee_ids
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Employees"

        headers = [
            "Employee ID",
            "First Name",
            "Middle Name",
            "Last Name",
            "Gender",
            "Date of Birth",
            "Phone Number",
            "Official Email",
            "Personal Email",
            "Marital Status",
            "Blood Group",
            "Nationality",
            "Father Name",
            "Mother Name",
            "Position",
            "Employee Type",
            "Attendance Category",
            "Employment Type",
            "Department",
            "Work Location",
            "Joining Date",
            "Salary",
            "PAN",
            "Aadhaar",
            "UAN",
            "PF Number",
            "Status",
        ]

        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="2563EB"
            )

        for employee in employees:

            worksheet.append([
                employee.employee_id,
                employee.first_name,
                employee.middle_name or "",
                employee.last_name,
                employee.gender or "",
                employee.date_of_birth or "",
                employee.phone_number,
                employee.official_email,
                employee.personal_email or "",
                employee.marital_status or "",
                employee.blood_group or "",
                employee.nationality,
                employee.father_name or "",
                employee.mother_name or "",
                employee.position,
                employee.employee_type or "",
                employee.attendance_category,
                employee.employment_type,
                (
                    employee.department.name
                    if employee.department
                    else ""
                ),
                employee.work_location or "",
                employee.joining_date,
                employee.salary,
                employee.pan,
                employee.aadhaar,
                employee.uan or "",
                employee.pf_number or "",
                "Active" if employee.is_active else "Inactive",
            ])

        for column in worksheet.columns:

            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 35)

        response = HttpResponse(
            content_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        response["Content-Disposition"] = (
            'attachment; filename="employees.xlsx"'
        )

        workbook.save(response)

        return response
        # ==========================================
    
    # ==========================================
    # IMPORT EMPLOYEES FROM EXCEL / CSV
    # ==========================================
    @action(
        detail=False,
        methods=["post"],
        url_path="import",
        parser_classes=[MultiPartParser, FormParser],
    )
    def import_employees(self, request):
        uploaded_file = request.FILES.get("file")

        if not uploaded_file:
            return Response(
                {"message": "Please upload an Excel or CSV file."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        filename = uploaded_file.name.lower()

        try:
            # Read the uploaded file
            if filename.endswith(".xlsx"):
                workbook = load_workbook(
                    uploaded_file,
                    read_only=True,
                    data_only=True,
                )
                worksheet = workbook["Employees"] if "Employees" in workbook.sheetnames else workbook.active
                rows = worksheet.iter_rows(values_only=True)

                headers = next(rows, None)

                if not headers:
                    workbook.close()
                    return Response(
                        {"message": "The Excel sheet is empty."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                headers = [
                    str(header).strip() if header is not None else ""
                    for header in headers
                ]

                records = [
                    dict(zip(headers, row))
                    for row in rows
                    if any(value is not None for value in row)
                ]
                workbook.close()

            elif filename.endswith(".csv"):
                text = uploaded_file.read().decode("utf-8-sig")
                records = list(csv.DictReader(io.StringIO(text)))

            elif filename.endswith(".xls"):
                return Response(
                    {
                        "message": (
                            "Legacy .xls files are not supported by this "
                            "endpoint yet. Save the workbook as .xlsx "
                            "and upload it again."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            else:
                return Response(
                    {"message": "Unsupported file format."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if not records:
                return Response(
                    {"message": "No employee records found in the file."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # Normalize Excel headers to Employee model/serializer fields.
            header_map = {
                "employee id": "employee_id",
                "first name": "first_name",
                "middle name": "middle_name",
                "last name": "last_name",
                "date of birth": "date_of_birth",
                "phone number": "phone_number",
                "official email": "official_email",
                "personal email": "personal_email",
                "marital status": "marital_status",
                "blood group": "blood_group",
                "father name": "father_name",
                "mother name": "mother_name",
                "employee type": "employee_type",
                "attendance category": "attendance_category",
                "employment type": "employment_type",
                "department": "department",
                "work location": "work_location",
                "joining date": "joining_date",
                "pf number": "pf_number",
                "status": "is_active",
            }

            created = 0
            failed = 0
            errors = []
            seen_emails = set()
            seen_ids = set()

            for row_number, record in enumerate(records, start=2):
                try:
                    data = {}

                    for header, value in record.items():
                        normalized_header = str(header).strip().lower()
                        field = header_map.get(
                            normalized_header,
                            normalized_header.replace(" ", "_"),
                        )

                        if value == "":
                            value = None

                        data[field] = value

                    # Normalize dates supplied as Excel date cells or strings.
                    for field in ("date_of_birth", "joining_date"):
                        value = data.get(field)

                        if isinstance(value, datetime):
                            data[field] = value.date()
                        elif isinstance(value, date):
                            pass
                        elif isinstance(value, str) and value.strip():
                            parsed = None
                            for fmt in (
                                "%Y-%m-%d",
                                "%d-%m-%Y",
                                "%d/%m/%Y",
                            ):
                                try:
                                    parsed = datetime.strptime(
                                        value.strip(), fmt
                                    ).date()
                                    break
                                except ValueError:
                                    continue

                            if parsed is None:
                                raise ValueError(
                                    f"Invalid {field} format: {value}"
                                )

                            data[field] = parsed

                    # Map the department name to its database record.
                    department_value = data.get("department")

                    if department_value not in (None, ""):
                        department = Department.objects.filter(
                            name__iexact=str(department_value).strip()
                        ).first()

                        if not department:
                            raise ValueError(
                                f"Department '{department_value}' "
                                "does not exist. Create it first."
                            )

                        data["department"] = department.pk
                    else:
                        data.pop("department", None)

                    # Normalize Active/Inactive values.
                    active_value = data.get("is_active")

                    if isinstance(active_value, str):
                        active_value = active_value.strip().lower()

                        if active_value in ("active", "true", "1", "yes"):
                            data["is_active"] = True
                        elif active_value in (
                            "inactive", "false", "0", "no"
                        ):
                            data["is_active"] = False

                    # Employee IDs are generated by your model.
                    # Do not attempt to override an auto-generated ID.
                    data.pop("employee_id", None)

                    # Avoid duplicate rows in the same import.
                    email = data.get("official_email")
                    if email:
                        email_key = str(email).strip().lower()

                        if email_key in seen_emails:
                            raise ValueError(
                                f"Duplicate official email in file: {email}"
                            )

                        seen_emails.add(email_key)

                    # Validate and save through your existing serializer.
                    serializer = self.get_serializer(data=data)
                    serializer.is_valid(raise_exception=True)
                    employee_login = CURRENT_USER.objects.filter(
                        email__iexact=str(data.get("official_email") or ""),
                        role="viewer",
                    ).first()
                    serializer.save(user=employee_login)

                    created += 1

                except (ValidationError, ValueError, TypeError) as exc:
                    failed += 1

                    if isinstance(exc, ValidationError):
                        detail = exc.detail
                    else:
                        detail = str(exc)

                    errors.append({
                        "row": row_number,
                        "error": detail,
                    })

                except Exception as exc:
                    failed += 1
                    errors.append({
                        "row": row_number,
                        "error": str(exc),
                    })

            return Response(
                {
                    "message": "Employee import completed.",
                    "filename": uploaded_file.name,
                    "total": len(records),
                    "created": created,
                    "failed": failed,
                    "errors": errors,
                },
                status=(
                    status.HTTP_200_OK
                    if failed == 0
                    else status.HTTP_207_MULTI_STATUS
                    if created > 0
                    else status.HTTP_400_BAD_REQUEST
                ),
            )

        except Exception as exc:
            return Response(
                {"message": f"Unable to process file: {str(exc)}"},
                status=status.HTTP_400_BAD_REQUEST,
            )


# ==========================================================
# Department ViewSet
# ==========================================================
class DepartmentViewSet(ModelViewSet):

    queryset = Department.objects.all().order_by("name")

    serializer_class = DepartmentSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminOrHRStrict
    ]
